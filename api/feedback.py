"""Reviewer feedback endpoints: false-positive flags and the correction chat."""
from __future__ import annotations

import asyncio
import base64
import json as _json
import os
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel

from canada_pipeline.region import get_region

router = APIRouter()
_REGION = get_region()
APP_REGION = _REGION.name      # label written to the feedback logs


# ─────────────────────────────────────────────
# False positives — validator flags a revision the checker wrongly called invalid
# ─────────────────────────────────────────────

class FalsePositiveRequest(BaseModel):
    session_id: str
    revision_key: str                 # stable per revision: "<paragraph_index>:<text hash>"
    revision_index: int = -1          # position in the Revisions list
    excel_row: int = -1               # 0-based data row in the validation Excel, if matched
    revision_type: str = ""
    revision_text: str = ""
    previous_context: str = ""
    paragraph_index: int = -1
    validator_reason: str = ""
    suggested_fix: str = ""
    severity: str = "medium"          # critical | medium | low
    note: str = ""
    emp_id: str = ""
    user_name: str = ""
    doc_name: str = ""


def _mark_false_positive_in_excel(excel_path: str, excel_row: int, severity: str, note: str) -> None:
    """Write/clear `false_positive` + `fp_note` columns on the matched row of the
    session's validation Excel (same convention as the `status` column)."""
    import openpyxl
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    if not ws or excel_row < 0:
        return
    header = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    cols = {}
    for name in ("false_positive", "fp_note"):
        col = next((i for i, h in enumerate(header, 1) if h and str(h).strip().lower() == name), None)
        if col is None:
            col = ws.max_column + 1
            ws.cell(row=1, column=col, value=name)
            header = list(header) + [name]
        cols[name] = col
    ws.cell(row=excel_row + 2, column=cols["false_positive"], value=severity or None)
    ws.cell(row=excel_row + 2, column=cols["fp_note"], value=(note or None))
    wb.save(excel_path)


def _refresh_session_fp_files(session_id: str, excel_row: int = -1, severity: str = "", note: str = ""):
    """Per-job copy: regenerate false_positives.xlsx in the session's S3 folder and mark the
    matched row of that session's validation Excel."""
    from src import false_positive_log as fp, session_store
    meta = session_store.load_meta(session_id)
    if meta is None:
        return
    fp.write_session_workbook(session_id)
    try:
        excel_name = meta.get("excel_filename")
        if excel_name and excel_row >= 0:
            session_store.edit_file(
                session_id, excel_name,
                lambda path: _mark_false_positive_in_excel(path, excel_row, severity, note))
    except Exception as e:
        print(f"[FALSE-POSITIVE] Could not mark validation Excel: {e}")


@router.post("/api/validation/false-positive")
async def flag_false_positive(request: FalsePositiveRequest):
    """Record (or update) a false-positive flag with a severity."""
    from src import false_positive_log as fp
    severity = (request.severity or "medium").lower()
    if severity not in fp.SEVERITIES:
        raise HTTPException(400, f"severity must be one of {', '.join(fp.SEVERITIES)}")
    row = {**request.model_dump(), "severity": severity, "region": APP_REGION}
    stored = await asyncio.to_thread(fp.upsert, row)
    await asyncio.to_thread(_refresh_session_fp_files, request.session_id, request.excel_row, severity, request.note)
    return {"status": "flagged", "severity": stored["severity"], "revision_key": request.revision_key}


@router.delete("/api/validation/false-positive")
async def unflag_false_positive(session_id: str, revision_key: str, emp_id: str = "", excel_row: int = -1):
    from src import false_positive_log as fp
    removed = await asyncio.to_thread(fp.remove, session_id, revision_key, emp_id)
    await asyncio.to_thread(_refresh_session_fp_files, session_id, excel_row, "", "")
    return {"status": "removed" if removed else "not_found"}


@router.get("/api/validation/false-positives")
async def list_false_positives(session_id: str):
    """Flags already made on this session (so the viewer can show them on reload)."""
    from src import false_positive_log as fp
    rows = await asyncio.to_thread(fp.session_flags, session_id)
    return {"flags": [
        {"revision_key": r["revision_key"], "severity": r["severity"], "note": r.get("note", ""), "emp_id": r.get("emp_id", "")}
        for r in rows
    ]}



# ─────────────────────────────────────────────
# Correction chat — retag a selected XML block from a plain-language instruction
# ─────────────────────────────────────────────

class CorrectionRequest(BaseModel):
    block_xml: str
    instruction: str
    context_before: str = ""
    context_after: str = ""
    parent_tag: str = ""
    pxml_head: str = ""               # start of the document (carries the DOCTYPE)
    dtd_family: str = "auto"
    history: List[dict] = []          # earlier turns: [{role: user|assistant, content}]
    session_id: str = ""
    emp_id: str = ""
    user_name: str = ""
    doc_name: str = ""


async def _parse_correction_payload(request: Request):
    """Accept either JSON or multipart/form-data. Multipart allows screenshot uploads."""
    ctype = request.headers.get("content-type", "").lower()
    if "multipart/form-data" in ctype:
        form = await request.form()

        def _field(name: str, default: str = "") -> str:
            value = form.get(name)
            return value if isinstance(value, str) else default

        history_raw = _field("history", "[]")
        try:
            history = _json.loads(history_raw) if history_raw else []
        except Exception:
            history = []

        images = []
        for file_field in list(form.getlist("images")) + list(form.getlist("image")):
            if not getattr(file_field, "filename", None):
                continue
            payload = await file_field.read()
            images.append({
                "filename": file_field.filename or "",
                "media_type": file_field.content_type or "image/png",
                "data": base64.b64encode(payload).decode("utf-8"),
            })

        return {
            "block_xml": _field("block_xml", ""),
            "instruction": _field("instruction", ""),
            "context_before": _field("context_before", ""),
            "context_after": _field("context_after", ""),
            "parent_tag": _field("parent_tag", ""),
            "pxml_head": _field("pxml_head", ""),
            "dtd_family": _field("dtd_family", "auto"),
            "history": history,
            "session_id": _field("session_id", ""),
            "emp_id": _field("emp_id", ""),
            "user_name": _field("user_name", ""),
            "doc_name": _field("doc_name", ""),
            "images": images,
        }

    payload = await request.json()
    if not isinstance(payload, dict):
        payload = {}
    history = payload.get("history") or []
    images = payload.get("images") or payload.get("image_data") or []
    if isinstance(images, dict):
        images = [images]
    return {
        "block_xml": payload.get("block_xml", ""),
        "instruction": payload.get("instruction", ""),
        "context_before": payload.get("context_before", ""),
        "context_after": payload.get("context_after", ""),
        "parent_tag": payload.get("parent_tag", ""),
        "pxml_head": payload.get("pxml_head", ""),
        "dtd_family": payload.get("dtd_family", "auto"),
        "history": history,
        "session_id": payload.get("session_id", ""),
        "emp_id": payload.get("emp_id", ""),
        "user_name": payload.get("user_name", ""),
        "doc_name": payload.get("doc_name", ""),
        "images": images,
    }


@router.post("/api/correction/propose")
async def propose_correction_endpoint(request: Request):
    """Propose a corrected version of the selected block. The proposal is checked against
    the DTD before it is returned (see canada_pipeline/correction_chat.py). Supports either
    JSON payloads or multipart image uploads for the AI-to-XML review flow."""
    from canada_pipeline.correction_chat import propose_correction
    from src import correction_log
    import uuid

    payload = await _parse_correction_payload(request)
    try:
        result = await asyncio.to_thread(
            propose_correction,
            block_xml=payload.get("block_xml", ""), instruction=payload.get("instruction", ""),
            context_before=payload.get("context_before", ""), context_after=payload.get("context_after", ""),
            parent_tag=payload.get("parent_tag", ""), pxml_head=payload.get("pxml_head", ""),
            dtd_family=payload.get("dtd_family", "auto"), history=payload.get("history") or [],
            images=payload.get("images") or [],
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        print(f"[CORRECTION] failed: {e}")
        raise HTTPException(503, f"The correction assistant is unavailable: {e}")
    proposal_id = uuid.uuid4().hex[:12]
    await asyncio.to_thread(correction_log.record_proposal, {
        "proposal_id": proposal_id, "session_id": payload.get("session_id", ""), "doc_name": payload.get("doc_name", ""),
        "emp_id": payload.get("emp_id", ""), "user_name": payload.get("user_name", ""),
        "dtd_family": payload.get("dtd_family", "auto"), "parent_tag": payload.get("parent_tag", ""),
        "instruction": payload.get("instruction", ""), "original_xml": payload.get("block_xml", ""),
        "proposed_xml": result["xml"], "explanation": result["explanation"],
        "valid": "yes" if result["valid"] else "no",
        "issues": "; ".join(i["message"] for i in result["issues"]),
        "model": result["model"], "region": APP_REGION,
    })
    return {**result, "proposal_id": proposal_id}


class CorrectionAppliedRequest(BaseModel):
    proposal_id: str
    applied: bool = True


@router.post("/api/correction/applied")
async def correction_applied(request: CorrectionAppliedRequest):
    from src import correction_log
    ok = await asyncio.to_thread(correction_log.mark_applied, request.proposal_id, request.applied)
    return {"status": "updated" if ok else "not_found"}


