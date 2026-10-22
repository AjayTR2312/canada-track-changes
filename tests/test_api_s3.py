"""End-to-end API tests against a fake S3: the full job / validation-session / batch lifecycles, and
the property the migration exists for -- a *different pod* (no local memory, no local disk) can
answer for state created by another one."""
import io
import os
import shutil
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient


def _zip_bytes(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _xlsx_bytes(rows) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _new_pod():
    """Forget everything a pod would hold in memory and on its scratch disk."""
    import pipeline_api as api
    from src import job_store
    job_store.jobs._local.clear()
    job_store.batches._local.clear()
    shutil.rmtree(api._JOBS_ROOT, ignore_errors=True)
    shutil.rmtree(os.path.join(os.path.dirname(api._JOBS_ROOT), "pxml_pipeline_batches"), ignore_errors=True)


@pytest.fixture()
def api(s3, monkeypatch):
    import pipeline_api
    monkeypatch.setattr(pipeline_api, "executor", ThreadPoolExecutor(max_workers=2))

    def fake_pipeline(job_id, docx_path, pxml_path, output_dir, product_type, images_dir=None, processing_mode="chapter"):
        pipeline_api._write_job_status(output_dir, 1, "Track Changes Extraction", 10)
        with open(os.path.join(output_dir, "pipeline_log.txt"), "w") as f:
            f.write(f"log of {job_id}\n")
        out = os.path.join(output_dir, "chap_20260101_000000_final.pxml")
        with open(out, "w") as f:
            f.write("<pxml/>")
        pipeline_api._write_job_status(output_dir, 11, "Post Processing", 95)
        pipeline_api._upload_outputs_with_retry(job_id, output_dir)
        return {"status": "completed", "output_file": out, "job_id": job_id, "llm_calls": 3}

    def fake_validation(job_id, output_dir, output_file):
        path = os.path.join(output_dir, "final_validation_report_20260101.xlsx")
        with open(path, "wb") as f:
            f.write(_xlsx_bytes([["a", "b"], [1, 2]]))
        pipeline_api._upload_outputs_with_retry(job_id, output_dir)
        return path

    monkeypatch.setattr(pipeline_api, "_run_pipeline_job", fake_pipeline)
    monkeypatch.setattr(pipeline_api, "_run_validation_sync", fake_validation)
    _new_pod()
    with TestClient(pipeline_api.app) as client:
        yield client
    _new_pod()


def _wait(predicate, timeout=20.0):
    end = time.time() + timeout
    while time.time() < end:
        v = predicate()
        if v:
            return v
        time.sleep(0.1)
    raise AssertionError("timed out")


# ─────────────────────────────── health ───────────────────────────────
def test_health_reports_storage(api):
    j = api.get("/api/health").json()
    assert j["status"] == "healthy" and j["storage"]["ok"] is True


# ─────────────────────────────── single job lifecycle ───────────────────────────────
def test_job_lifecycle_survives_a_pod_change(api, s3):
    r = api.post("/api/process",
                 files={"docx_file": ("chap.docx", b"DOCX"), "pxml_file": ("chap.pxml", b"<x/>"),
                        "images_zip": ("imgs.zip", _zip_bytes({"p1.png": "img", "chap.pdf": "pdf"}))},
                 data={"product_type": "looseleaf", "user_email": "a@x.com", "emp_id": "42"})
    assert r.status_code == 200
    job_id = r.json()["job_id"]
    # inputs are in S3 from the moment the job exists
    assert s3.get_bytes(f"uploads/{job_id}/chap.docx") == b"DOCX"
    assert s3.get_bytes(f"uploads/{job_id}/images/p1.png") == b"img"

    _wait(lambda: api.get(f"/api/status/{job_id}").json()["status"] == "completed")
    assert s3.get_json(f"jobs/{job_id}/job.json")["status"] == "completed"

    _new_pod()        # ---- a different pod (or this one after a restart) answers from here on ----
    st = api.get(f"/api/status/{job_id}").json()
    assert st["status"] == "completed" and st["progress_pct"] == 100 and "_future" not in st

    z = api.get(f"/api/download/{job_id}")
    assert z.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert "chap_20260101_000000_final.pxml" in names and "final_validation_report_20260101.xlsx" in names
    assert not any(n.startswith("_") for n in names)

    v = api.get(f"/api/download-validation/{job_id}")
    assert v.status_code == 200 and v.headers["content-disposition"].startswith("attachment")
    assert api.get(f"/api/log/{job_id}").json()["log"].startswith("log of")
    assert {f["filename"] for f in api.get(f"/api/list-outputs/{job_id}").json()["files"]} >= {"pipeline_log.txt"}
    assert api.get(f"/api/download-by-session/{job_id}").status_code == 200
    assert api.get("/api/status/nope").status_code == 404

    # usage was logged to S3 and the recent-jobs view is built from S3 only
    recent = api.get("/api/auth/recent-jobs", params={"emp_id": "42"}).json()
    row = next(r for r in recent if r["session_id"] == job_id)
    assert row["run_outcome"] == "success" and row["download_available"] and row["validation_done"]
    assert row["file_name"] == "chap.docx"
    logs = api.get("/api/auth/usage-logs", params={"emp_id": "42"}).json()
    assert logs and logs[0]["llm_calls"] == 3


def test_revalidation_on_another_pod_restores_files_from_s3(api, s3, monkeypatch):
    import pipeline_api
    r = api.post("/api/process",
                 files={"docx_file": ("chap.docx", b"DOCX"), "pxml_file": ("chap.pxml", b"<x/>"),
                        "images_zip": ("imgs.zip", _zip_bytes({"p1.png": "img"}))},
                 data={"emp_id": "42"})
    job_id = r.json()["job_id"]
    _wait(lambda: api.get(f"/api/status/{job_id}").json()["status"] == "completed")

    _new_pod()                                    # this pod has no files and no memory of the job
    assert not os.path.exists(os.path.join(pipeline_api._JOBS_ROOT, job_id))
    seen = {}

    def fake_validation(jid, output_dir, output_file):
        job_dir = os.path.dirname(output_dir)
        seen.update(pxml=os.path.isfile(output_file),
                    docx=os.path.isfile(os.path.join(job_dir, "chap.docx")),
                    images=os.path.isfile(os.path.join(job_dir, "revision_images", "p1.png")))
        path = os.path.join(output_dir, "final_validation_report_20270101.xlsx")
        with open(path, "wb") as f:
            f.write(_xlsx_bytes([["new"]]))
        pipeline_api._upload_outputs_with_retry(jid, output_dir)
        return path

    monkeypatch.setattr(pipeline_api, "_run_validation_sync", fake_validation)
    assert api.post(f"/api/validate/{job_id}").status_code == 200
    assert seen == {"pxml": True, "docx": True, "images": True}                 # restored from S3 on demand
    assert s3.object_exists(f"outputs/{job_id}/final_validation_report_20270101.xlsx")
    assert api.get(f"/api/status/{job_id}").json()["current_step_name"] == "Completed"


def test_cancel_works_across_pods(api, s3):
    import pipeline_api
    from src import job_store
    s3.put_json("jobs/xyz/job.json", {"job_id": "xyz", "status": "running", "started_at": time.time(),
                                      "output_dir": "/other/pod/xyz/output"})
    s3.put_json("jobs/xyz/heartbeat.json", {"at": time.time()})
    r = api.post("/api/cancel/xyz").json()
    assert r["status"] == "cancelling" and job_store.is_cancelled("xyz")
    # the worker (on whichever pod runs it) notices at its next step boundary
    out = os.path.join(pipeline_api._JOBS_ROOT, "xyz", "output")
    os.makedirs(out, exist_ok=True)
    with pytest.raises(pipeline_api.PipelineCancelled):
        pipeline_api._write_job_status(out, 2, "Section Deletion", 20)


def test_stopped_pod_job_is_reported_interrupted(api, s3):
    s3.put_json("jobs/old/job.json", {"job_id": "old", "status": "running", "started_at": time.time() - 3600})
    s3.put_json("jobs/old/heartbeat.json", {"at": time.time() - 3600})
    st = api.get("/api/status/old").json()
    assert st["status"] == "failed" and st["error_info"]["retryable"] is True


# ─────────────────────────────── validation sessions ───────────────────────────────
def test_validation_session_lifecycle(api, s3):
    r = api.post("/api/validation/create-session",
                 files={"docx_file": ("doc.docx", b"DOCXBYTES"), "pxml_file": ("doc.pxml", b"<x>1</x>"),
                        "excel_file": ("report.xlsx", _xlsx_bytes([["h1"], ["r0"], ["r1"]])),
                        "pdf_file": ("doc.pdf", b"%PDF-" + bytes(range(50)))},
                 data={"emp_id": "42", "session_name": "My doc"})
    assert r.status_code == 200
    meta = r.json()
    sid = meta["session_id"]
    assert s3.get_bytes(f"validation_sessions/{sid}/doc.docx") == b"DOCXBYTES"

    _new_pod()
    got = api.get(f"/api/validation/session/{sid}").json()
    assert got["session_name"] == "My doc" and got["storage_prefix"].endswith(f"validation_sessions/{sid}/")
    assert [s["session_id"] for s in api.get("/api/validation/sessions", params={"emp_id": "42"}).json()["sessions"]] == [sid]
    assert api.get("/api/validation/sessions", params={"emp_id": "7"}).json()["sessions"] == []

    # files stream out of S3, with Range support (the PDF viewer relies on it)
    full = api.get(f"/api/validation/session/{sid}/file/doc.pdf")
    assert full.status_code == 200 and full.headers["accept-ranges"] == "bytes" and len(full.content) == 55
    part = api.get(f"/api/validation/session/{sid}/file/doc.pdf", headers={"Range": "bytes=5-9"})
    assert part.status_code == 206 and part.content == bytes(range(5)) and part.headers["content-range"] == "bytes 5-9/55"
    assert api.get(f"/api/validation/session/{sid}/file/doc.pdf", headers={"Range": "bytes=999-"}).status_code == 416
    assert api.get(f"/api/validation/session/{sid}/file/ghost.docx").status_code == 404
    assert api.get("/api/validation/session/ghost").status_code == 404

    # progress updates r/m/w the record AND the Excel in S3
    api.post("/api/validation/set-total-revisions", data={"session_id": sid, "total_revisions": 2, "total_invalid": 2})
    u = api.post("/api/validation/update-revision-status", json={"session_id": sid, "revision_index": 0, "status": "complete"}).json()
    assert u["progress_text"] == "1/2" and u["current_revision_index"] == 1
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(s3.get_bytes(f"validation_sessions/{sid}/report.xlsx"))).active
    assert [c.value for c in ws[1]] == ["h1", "status"] and ws["B2"].value == "complete"
    api.post("/api/validation/update-position", data={"session_id": sid, "current_revision_index": 1})
    assert api.get(f"/api/validation/session/{sid}").json()["current_revision_index"] == 1

    # edited PXML is saved to S3
    assert api.post("/api/validation/save-pxml", json={"session_id": sid, "pxml_content": "<x>edited</x>"}).json()["status"] == "saved"
    assert s3.get_bytes(f"validation_sessions/{sid}/doc.pxml") == b"<x>edited</x>"
    # false-positive flags: log in S3 + per-session workbook + marked Excel row
    fp = {"session_id": sid, "revision_key": "0:abc", "excel_row": 0, "severity": "critical", "emp_id": "42", "revision_text": "t"}
    assert api.post("/api/validation/false-positive", json=fp).json()["status"] == "flagged"
    assert s3.object_exists(f"validation_sessions/{sid}/false_positives.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(s3.get_bytes(f"validation_sessions/{sid}/report.xlsx"))).active
    assert "false_positive" in [c.value for c in ws[1]]
    assert api.get("/api/validation/false-positives", params={"session_id": sid}).json()["flags"][0]["severity"] == "critical"
    assert api.delete("/api/validation/false-positive", params={"session_id": sid, "revision_key": "0:abc", "emp_id": "42", "excel_row": 0}).json()["status"] == "removed"
    assert not s3.object_exists(f"validation_sessions/{sid}/false_positives.xlsx")

    # extra files can be added to a session; the reserved progress-record name cannot be overwritten
    up = api.post(f"/api/validation/session/{sid}/upload-file", files={"file": ("Failed_revisions.xlsx", _xlsx_bytes([["x"]]))})
    assert up.json()["status"] == "uploaded"
    assert api.post(f"/api/validation/session/{sid}/upload-file", files={"file": ("_session.json", b"evil")}).status_code == 400

    # DTD validation can read the saved PXML straight from the session (no pxml_content posted)
    assert api.post("/api/validation/validate-dtd", json={"session_id": "ghost"}).status_code == 404

    # resume rebuilds the record from the files when the record itself is gone
    s3.delete_object(f"validation_sessions/{sid}/_session.json")
    s3.delete_object(f"validation_index/42/{sid}.json")
    res = api.post("/api/validation/resume", data={"session_id": sid, "emp_id": "42"}).json()
    assert res["docx_filename"] == "doc.docx" and res["pxml_filename"] == "doc.pxml" and res["pdf_filename"] == "doc.pdf"
    assert res["excel_filename"] == "report.xlsx" and res["failed_revisions_filename"] == "Failed_revisions.xlsx"
    assert api.post("/api/validation/resume-from-sharepoint", data={"session_id": sid}).status_code == 200     # legacy path
    assert api.post("/api/validation/resume", data={"session_id": "ghost"}).status_code == 404

    assert api.delete(f"/api/validation/session/{sid}").json()["status"] == "deleted"
    assert s3.list_objects(f"validation_sessions/{sid}/") == [] and api.get(f"/api/validation/session/{sid}").status_code == 404


def test_draft_revisions_are_shared_via_s3(api, s3):
    assert api.get("/api/draft-revisions").status_code == 404
    wb = _xlsx_bytes([["Row Index", "Section", "Reason", "Change Type"], [1, "§1", "no match", "inline"],
                      [2, "§2", "PLEASE VERIFY", "inline"]])
    assert api.post("/api/draft-revisions/upload", files={"file": ("Failed_revisions.xlsx", wb)}).status_code == 200
    _new_pod()
    j = api.get("/api/draft-revisions").json()
    assert j["total"] == 1 and j["rows"][0]["Section"] == "§1"


# ─────────────────────────────── users ───────────────────────────────
def test_auth_endpoints_use_s3(api, s3):
    assert api.post("/api/auth/login", json={"emp_id": "5"}).status_code == 404
    assert api.post("/api/auth/register", json={"emp_id": "5", "name": "Eve", "email": "e@x.com", "dept": "QA"}).status_code == 200
    assert s3.get_json("users/5.json")["name"] == "Eve"
    assert api.post("/api/auth/login", json={"emp_id": "5"}).json()["name"] == "Eve"
    api.put("/api/auth/update", json={"emp_id": "5", "name": "Eve 2", "email": "e@x.com", "dept": "QA"})
    assert api.post("/api/auth/login", json={"emp_id": "5"}).json()["name"] == "Eve 2"


# ─────────────────────────────── batches ───────────────────────────────
def _batch_files():
    return [("files", ("ch1.docx", b"d1")), ("files", ("ch1.xml", b"<x/>")),
            ("files", ("ch1.zip", _zip_bytes({"p.png": "i", "ch1.pdf": "pdf"})))]


def test_batch_lifecycle_across_pods(api, s3):
    plan = api.post("/api/batch/plan", files=_batch_files()).json()
    bid = plan["batch_id"]
    assert plan["groups"][0]["missing"] == [] and s3.object_exists(f"batches/{bid}/input/ch1.docx")

    _new_pod()                                    # another pod: the review table is restored from S3
    st = api.get(f"/api/batch/{bid}").json()
    assert st["status"] == "planned" and st["groups"][0]["docx"] == "ch1.docx"

    r = api.post(f"/api/batch/{bid}/start", data={"groups": __import__("json").dumps(plan["groups"]), "product_type": "looseleaf",
                                                    "emp_id": "42", "publication_name": "PUB 1"})
    assert r.status_code == 200 and r.json()["s3_output_folder"].startswith("PUB 1_Output_")
    _wait(lambda: api.get(f"/api/batch/{bid}").json()["status"] == "completed")

    st = api.get(f"/api/batch/{bid}").json()
    item = st["items"][0]
    assert item["status"] == "completed" and st["completed"] == 1 and item["progress_pct"] == 100
    assert item["s3_published"] >= 2 and item["s3_location"].startswith("s3://test-bucket/published/42/PUB 1_Output_")
    published = [o["key"] for o in s3.list_objects("published/42/")]
    assert any(k.endswith("/ch1/chap_20260101_000000_final.pxml") for k in published)
    assert any("Batch started" in e["message"] for e in st["log"])

    _new_pod()                                    # still fully readable, and retry works from S3 inputs
    assert api.get(f"/api/batch/{bid}").json()["items"][0]["status"] == "completed"
    assert api.post(f"/api/batch/{bid}/retry-push/ch1").json()["s3_error"] == ""
    assert api.post(f"/api/batch/{bid}/retry/ch1").json()["status"] == "running"
    _wait(lambda: api.get(f"/api/batch/{bid}").json()["status"] == "completed"
          and api.get(f"/api/batch/{bid}").json()["items"][0]["status"] == "completed")
    assert api.post(f"/api/batch/{bid}/retry/nope").status_code == 404


def test_batch_from_s3_folder_publishes_next_to_input(api, s3):
    s3.put_bytes("incoming/Pub9/ch1.docx", b"d")
    s3.put_bytes("incoming/Pub9/ch1.xml", b"<x/>")
    s3.put_bytes("incoming/Pub9/ch1.zip", _zip_bytes({"a.png": "i"}))
    assert api.post("/api/batch/plan-s3", data={"s3_folder": "jobs/"}).status_code == 400          # reserved prefix
    assert api.post("/api/batch/plan-s3", data={"s3_folder": "incoming/empty"}).status_code == 400
    plan = api.post("/api/batch/plan-s3", data={"s3_folder": "s3://test-bucket/incoming/Pub9"}).json()
    bid = plan["batch_id"]
    assert s3.object_exists(f"batches/{bid}/input/ch1.docx")        # own copy kept for retries
    import json
    api.post(f"/api/batch/{bid}/start", data={"groups": json.dumps(plan["groups"]), "publication_name": "P9"})
    _wait(lambda: api.get(f"/api/batch/{bid}").json()["status"] == "completed")
    out = api.get(f"/api/batch/{bid}").json()
    assert out["source"] == "s3" and out["s3_output_folder"].startswith("P9_Output_")
    assert any(o["key"].startswith(f"incoming/{out['s3_output_folder']}/ch1/") for o in s3.list_objects("incoming/"))


def test_batch_cancel_flag_is_shared(api, s3):
    plan = api.post("/api/batch/plan", files=_batch_files()).json()
    bid = plan["batch_id"]
    assert api.post(f"/api/batch/{bid}/cancel").json()["status"] == "cancelling"
    assert s3.object_exists(f"batches/{bid}/cancel.flag")
    assert api.post("/api/batch/nope/cancel").status_code == 404
