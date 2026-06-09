"""
FastAPI backend for PXML Document Processor Pipeline.
Uses Celery + Redis for async job execution across multiple workers/pods.
Job state persisted in Redis (survives pod restarts).
"""
import os
import sys
import time
import uuid
import shutil
import zipfile
import threading
from dotenv import load_dotenv
load_dotenv()
from pathlib import Path
from typing import Optional
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware

# Pipeline imports (must be at module level for wildcard imports)
from continous_ncont_insertions import *
from newline_insertions import *
from insertions_paragraphs import *
from section_name_edge_cases import *

# Redis-backed job store (replaces in-memory dict)
from job_store import job_store, JobStatus

# Celery task imports
from celery_app import run_pipeline_task, run_validation_task

app = FastAPI(title="PXML Document Processor API", version="2.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


# ─────────────────────────────────────────────
# Startup: Clean up orphaned jobs from previous pod lifecycle
# ─────────────────────────────────────────────
@app.on_event("startup")
async def startup_cleanup():
    cleaned = job_store.mark_stale_jobs_failed()
    if cleaned:
        print(f"[STARTUP] Marked {cleaned} orphaned jobs as failed (pod restart recovery)")


PIPELINE_STEP_NAMES = {
    1: "Track Changes Extraction",
    2: "Section Deletion",
    3: "Section Number Replacement",
    4: "Section Insertion & Edge Cases",
    5: "Paragraph-Level Inline Changes",
    6: "Newline Paragraph Insertions",
    7: "New Footnote Insertions",
    8: "Failed Revisions Report",
    9: "Retry Failed Deletions",
    10: "Post Processing",
}


# ─────────────────────────────────────────────
# Pipeline runner (called by Celery worker OR locally)
# ─────────────────────────────────────────────
def _run_pipeline_logic(job_id: str, docx_path: str, pxml_path: str, output_dir: str, product_type: str, images_dir: str = None, output_file: str = None):
    """Execute the full pipeline, updating job status in Redis as it progresses."""
    job = job_store.get(job_id)
    if not job:
        raise RuntimeError(f"Job {job_id} not found in store")
    job.status = "running"
    job.started_at = time.time()

    os.makedirs(output_dir, exist_ok=True)
    pxml_basename = os.path.splitext(os.path.basename(pxml_path))[0]
    timestamp = time.strftime('%Y%m%d_%H%M%S')
    if not output_file:
        output_file = os.path.join(output_dir, f"{pxml_basename}_{timestamp}_final.pxml")
    missing_revisions_output = os.path.join(output_dir, "missing_revisions.txt")
    log_file = os.path.join(output_dir, "pipeline_log.txt")
    failed_log_path = os.path.join(output_dir, "failed_items.txt")

    job.output_dir = output_dir
    job.output_file = output_file
    job_store.update(job)

    # Capture logs
    from io import StringIO
    log_buffer = StringIO()

    class LogCapture:
        def __init__(self):
            self.original = sys.stdout

        def write(self, data):
            log_buffer.write(data)
            job.log = log_buffer.getvalue()[-50000:]  # Keep last 50k chars
            self.original.write(data)
            # Detect step transitions
            self._detect_step(data)

        def flush(self):
            self.original.flush()

        def _detect_step(self, text):
            text_lower = text.lower()
            markers = {
                1: ["pipeline-1", "pipeline started"],
                2: ["pipeline-2", "section deletion"],
                3: ["pipeline-3", "section (number) replacement"],
                4: ["pipeline-4", "section insertion"],
                5: ["pipeline 5", "paragraph-level inline"],
                6: ["pipeline 6", "newline paragraph"],
                7: ["pipeline 7", "footnote insertion"],
                8: ["consolidated failed", "failed revisions excel"],
                9: ["pipeline 9", "retry failed"],
                10: ["pipeline-10", "post processing"],
            }
            for step_id, step_markers in markers.items():
                for marker in step_markers:
                    if marker in text_lower:
                        job.current_step = step_id
                        job.current_step_name = PIPELINE_STEP_NAMES[step_id]
                        job.progress_pct = int((step_id / 10) * 100)
                        job.log = log_buffer.getvalue()[-50000:]
                        job_store.update(job)
                        return

    log_capture = LogCapture()
    log_f = open(log_file, 'w', encoding='utf-8')

    class Tee:
        def __init__(self, *streams):
            self.streams = streams
        def write(self, data):
            for s in self.streams:
                try:
                    s.write(data)
                except:
                    pass
        def flush(self):
            for s in self.streams:
                try:
                    s.flush()
                except:
                    pass

    sys.stdout = Tee(log_capture, log_f)
    sys.stderr = Tee(log_capture, log_f)

    try:
        # Import all pipeline dependencies
        from finaloutputs.chapter_section import extract_designator_pattern_using_llm
        from full_section_insertion_prevendor import PXMLInserter
        from inline_replacement_section import SectionLineReplacement
        from new_footnotes_insertions import process_all_footnote_insertions
        from section_insertion import SectionExtractor
        from track_changes_df import process_docx
        from section_deletion import section_deletion_filling
        from config import API_KEY, MODEL_OPUS, MODEL_SONNET, build_anthropic_client as _build_client_config
        from section_split_ import handle_section_split
        from postprocessing import postprocessing_fortags
        from paragraph_level_trackchanges import parse_pxml_sections_from_str_new, process_with_llm
        from footnotes_level_trackchanges import process_with_llm_footnotes
        from inline_footnotes_insertion import (
            process_all_footnotes_with_llm, parse_pxml_sections_from_str_dict_positions, build_anthropic_client
        )
        from newline_insertions_new import return_closest_previous_par, process_all_newline_insertions_with_llm
        from sectionnumber_filling import fill_sections_backward_verbose
        from revisions_images import extract_page_texts
        from failed_revisions_handler import retry_failed_inline_deletions
        import pandas as pd
        import re
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl import load_workbook

        PRODUCT_TYPE = product_type

        # ═══ PIPELINE 1: Track Changes → DataFrame ═══
        start_time = time.time()
        print(f"\n{'='*80}")
        print(f"Pipeline started at: {time.strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"{'='*80}")

        results = process_docx(docx_path)
        extractor_obj = results['extractor']
        result_with_ba = extractor_obj.build_before_after_df()
        try:
            result_with_ba = extractor_obj.enrich_footnote_anchors()
        except RuntimeError as e:
            print(f"  [WARN] enrich_footnote_anchors skipped: {e}")

        no_changes_df = results['no_changes_df']
        df_remaining = results['df_remaining']
        result = results['result']
        df_changes = results['df_changes']
        chapter_number = results['chapter_number']
        print(f"\nChapter from DOCX: {chapter_number}")

        # Superscript extraction
        print("\n[SUPERSCRIPT] Extracting footnote superscript numbers from PDF...")
        from superscript_extraction_from_pdf import convert_docx_to_pdf, process_excel as process_superscripts
        try:
            pdf_path_for_superscript = convert_docx_to_pdf(docx_path, output_dir=output_dir)
            superscript_excel = os.path.join(output_dir, "superscript_extraction.xlsx")
            _, result_with_ba = process_superscripts(
                docx_path=docx_path, pdf_path=pdf_path_for_superscript,
                output_excel_path=superscript_excel, result_with_ba=result_with_ba
            )
        except Exception as e:
            print(f"[SUPERSCRIPT] WARNING: {e}")

        inline_insertion_df = result[result['type'] == 'inline_insertion']

        with open(pxml_path, 'r', encoding='utf-8') as f:
            pxml_content = f.read()

        pxml_chapter_number, section_pattern, designator_patterns = extract_designator_pattern_using_llm(pxml_content)
        separator = ':'
        if designator_patterns:
            sep_match = re.search(r'[^0-9A-Za-z]', designator_patterns[0])
            if sep_match:
                separator = re.escape(sep_match.group())

        if chapter_number is None or (isinstance(chapter_number, str) and chapter_number.strip() == ''):
            chapter_number = pxml_chapter_number

        # ═══ PIPELINE 2: Section Deletion ═══
        print(f"\n{'='*80}\nPIPELINE-2: Section Deletion\n{'='*80}")
        with open(pxml_path, 'r', encoding='utf-8') as f:
            pxml_content = f.read()
        updated_pxmlcontent, remaining_rows_inlinedeletion, rows_to_delete = \
            section_deletion_filling(result, df_changes, chapter_number, no_changes_df, df_remaining, pxml_content, verbose=True)

        # ═══ PIPELINE 3: Section Number Replacement ═══
        print(f"\n{'='*80}\nPIPELINE-3: Section (Number) Replacement\n{'='*80}")
        inlinereplacement_df = result_with_ba[(result_with_ba['type'] == 'inline_replacement') & (result_with_ba['is_footnote_x'] == False)]
        processor = SectionLineReplacement(inlinereplacement_df, API_KEY, output_file)
        processed_df = processor.process_dataframe(df_changes)
        updated_pxmlcontent, updated_inlinereplacement_df, replaced_indices = \
            processor.create_deletion_column(chapter_number, updated_pxmlcontent, missing_revisions_output)

        # ═══ PIPELINE 4: Section Insertion & Edge Cases ═══
        print(f"\n{'='*80}\nPIPELINE-4: Section Insertion & Edge Cases\n{'='*80}")
        sec_extractor = SectionExtractor(pxml_path)
        inline_insertion_df_without_footnotes = result_with_ba[
            (result_with_ba['type'] == 'inline_insertion') & (result_with_ba['is_footnote_x'] == False)
        ].drop_duplicates(subset=['idx'], keep='first')
        sections = sec_extractor.extract_inserted_sections(inline_insertion_df_without_footnotes, chapter_number)

        edgecases_in_section_insertion_df = remaining_insertions_df(inline_insertion_df_without_footnotes, sections)
        replacement_para_indices = set(result_with_ba[result_with_ba['type'] == 'inline_replacement']['idx'].tolist())
        edgecases_in_section_insertion_df = edgecases_in_section_insertion_df[
            ~edgecases_in_section_insertion_df['idx'].isin(replacement_para_indices)
        ]
        new_sections = create_edgecase_section_dict(edgecases_in_section_insertion_df, chapter_number, result_with_ba=result_with_ba)
        sections.extend(new_sections)

        new_sections = [s for s in new_sections if re.match(r'^(?:§\s*)?\d+:\d+', s['section_header'])]
        edgecase_processed_indexes = []
        if len(new_sections) == 1:
            for con in new_sections[0]['content']:
                if con.get('text') and re.match(r'^(?:§\s*)?\d+:\d+', con['text']):
                    section_number, section_text = extract_sectionedgecase_key(con['text'])
                    updated_pxmlcontent = update_designator_by_name_match(updated_pxmlcontent, section_text, section_number)
                    edgecase_processed_indexes.append(int(con['paragraph_index']))
                    break

        xml_context = sec_extractor.getting_contextxml()
        sections_insertion_indexes = list(edgecase_processed_indexes)

        pxml_chapter_number_p4, section_pattern_p4, designator_patterns_p4 = extract_designator_pattern_using_llm(updated_pxmlcontent)
        section_dict_p4 = parse_pxml_sections_from_str_new(updated_pxmlcontent, chapter_number)
        _client_p4 = _build_client_config()

        if len(sections) > 0:
            for each_section in sections:
                if each_section['start_df_index'] == each_section['end_df_index']:
                    continue
                header_desig = re.sub(r'^[§\s]+', '', str(each_section['section_header'])).strip()
                header_desig = header_desig.split()[0] if header_desig else ''

                prev_desig = None
                is_section_split = False
                previous_section_xml = None

                if header_desig and designator_patterns_p4:
                    keys_list_str = "\n".join(f"  - '{k}'" for k in designator_patterns_p4)
                    prompt = (
                        f"Given a NEW section designator '{header_desig}' that is being inserted,\n"
                        f"which existing designator from the list below would come immediately BEFORE it "
                        f"in logical/numerical order?\n\nAvailable designators (in order):\n{keys_list_str}\n\n"
                        f"Rules:\n- '14:11A' comes after '14:11' but before '14:12'\n"
                        f"- Reply with ONLY the key exactly as shown above, or 'NONE' if none comes before it."
                    )
                    try:
                        response = _client_p4.messages.create(model=MODEL_SONNET, max_tokens=30, temperature=0, messages=[{"role": "user", "content": prompt}])
                        answer = response.content[0].text.strip().strip("'\"")
                        if answer.upper() != "NONE" and answer in designator_patterns_p4:
                            prev_desig = answer
                    except:
                        pass

                if prev_desig and prev_desig in section_dict_p4:
                    previous_section_xml = section_dict_p4[prev_desig]
                    start_pi_check = int(each_section['start_paragraph_index'])
                    desig_pattern_re = re.compile(r'^§?\s*\d+' + separator + r'\d+', re.IGNORECASE)
                    rows_after = result_with_ba[result_with_ba['idx'] > start_pi_check].drop_duplicates(subset=['idx'], keep='first').sort_values('idx')
                    checked_rows = []
                    for _, row in rows_after.iterrows():
                        row_text = str(row.get('text', '')).strip()
                        if row_text and desig_pattern_re.match(row_text):
                            break
                        has_change = bool(row.get('has_changes', False))
                        if not has_change and not row_text:
                            continue
                        checked_rows.append((row['idx'], has_change, row_text[:80]))

                    if checked_rows and not all(hc for _, hc, _ in checked_rows):
                        is_section_split = True
                        updated_pxmlcontent, split_success = handle_section_split(
                            each_section=each_section, previous_section_xml=previous_section_xml,
                            result_with_ba=result_with_ba, sec_extractor=sec_extractor,
                            xml_context=xml_context, updated_pxmlcontent=updated_pxmlcontent,
                            client=_client_p4, model_name=MODEL_OPUS, output_file=output_file,
                        )
                        if split_success:
                            start_pi = int(each_section['start_paragraph_index'])
                            end_pi = int(each_section['end_paragraph_index'])
                            sections_insertion_indexes.extend(range(start_pi, end_pi + 1))

                if not is_section_split:
                    start_pi = int(each_section['start_paragraph_index'])
                    end_pi = int(each_section['end_paragraph_index'])
                    sections_insertion_indexes.extend(range(start_pi, end_pi + 1))
                    plaintext_tobe_converted_xml = each_section['full_text']
                    bullet_points = each_section.get('bullet_points', [])
                    tagged_xml = sec_extractor.convert_text_to_xml_tags(
                        xml_context, plaintext_tobe_converted_xml,
                        bullet_points=bullet_points, section_header=each_section['section_header'],
                    )
                    cleanedup_xml, designator = sec_extractor.cleanup_xml_fences(tagged_xml, fallback_designator=each_section['section_header'])
                    if designator:
                        designator = designator.replace('§', '').strip()
                    insertionsection_obj = PXMLInserter(output_file, updated_pxmlcontent)
                    result_pxml, designator_num = insertionsection_obj.insert_xml_by_designator(designator, cleanedup_xml, write_to_file=False)
                    if result_pxml is not None:
                        updated_pxmlcontent = result_pxml

        # ═══ PIPELINE 5: Paragraph-level inline changes ═══
        print("\n[PIPELINE 5] Paragraph-level inline changes...")
        inline_replacement_df = result[result['type'] == 'inline_replacement']
        inline_insertion_df = result[result['type'] == 'inline_insertion']
        inline_deletion_df = result[result['type'] == 'inline_deletion']

        pxml_chapter_number, section_pattern, designator_patterns = extract_designator_pattern_using_llm(updated_pxmlcontent)
        section_dict = parse_pxml_sections_from_str_new(updated_pxmlcontent, chapter_number)

        WORKSPACE_ID = os.getenv("TR_WORKSPACE_ID", "")
        MODEL_NAME = MODEL_OPUS
        client, model_name = build_anthropic_client(WORKSPACE_ID, MODEL_NAME)

        if 'text_before' in result_with_ba.columns:
            text_before_mask = result_with_ba['text_before'].notna()
            result_with_ba.loc[text_before_mask, 'text_before'] = (
                result_with_ba.loc[text_before_mask, 'text_before'].astype(str).str.replace(r'^\s*[\.,]+\s*$', '', regex=True)
            )

        result_with_ba = result_with_ba.drop_duplicates()
        already_processed = set(rows_to_delete['paragraph_index'].tolist()) | set(replaced_indices) | set(sections_insertion_indexes)
        result_with_ba_new = result_with_ba[~result_with_ba['idx'].isin(already_processed)]

        all_revisions_except_newline_insertions_new = result_with_ba_new
        result_with_ba_sorted = result_with_ba_new.sort_values('idx').reset_index(drop=True)
        valid_text = result_with_ba_sorted[
            result_with_ba_sorted['text_before'].notna() & (result_with_ba_sorted['text_before'] != '') & (result_with_ba_sorted['text_before'].str.len() > 5)
        ][['idx', 'text_before']].reset_index(drop=True)

        def get_previous_para(idx):
            candidates = valid_text[valid_text['idx'] < idx]
            return None if candidates.empty else candidates.iloc[-1]['text_before']

        def get_post_para(idx):
            candidates = valid_text[valid_text['idx'] > idx]
            return None if candidates.empty else candidates.iloc[0]['text_before']

        all_revisions_except_newline_insertions_new['previous_para'] = all_revisions_except_newline_insertions_new['idx'].apply(get_previous_para)
        all_revisions_except_newline_insertions_new['post_para'] = all_revisions_except_newline_insertions_new['idx'].apply(get_post_para)

        all_revisions_except_newline_insertions1 = all_revisions_except_newline_insertions_new[
            ((all_revisions_except_newline_insertions_new['text_before'] != '') | (all_revisions_except_newline_insertions_new['type'] == 'inline_deletion')) &
            (all_revisions_except_newline_insertions_new['has_changes'] == True) & (all_revisions_except_newline_insertions_new['is_footnote_x'] == False)
        ]
        all_revisions_except_newline_insertions1 = fill_sections_backward_verbose(
            all_revisions_except_newline_insertions1, no_changes_df, inline_replacement_df, inline_insertion_df, chapter_number,
            verbose=True, client=client, model_name=model_name, designator_patterns=designator_patterns
        )

        # Page images — load from pre-generated images folder (uploaded by user)
        page_texts, image_paths = None, None
        try:
            if images_dir and os.path.isdir(images_dir):
                print(f"[PAGE IMAGES] Loading pre-generated images from: {images_dir}")
                # Load page_texts from JSON if available
                page_texts_json = os.path.join(images_dir, "page_texts.json")
                if os.path.isfile(page_texts_json):
                    import json as _json
                    with open(page_texts_json, 'r', encoding='utf-8') as _f:
                        _raw = _json.load(_f)
                        page_texts = {int(k): v for k, v in _raw.items()}
                else:
                    # Fallback: try extracting from PDF in the images folder
                    pdf_files = [f for f in os.listdir(images_dir) if f.endswith('.pdf')]
                    if pdf_files:
                        from revisions_images import extract_page_texts
                        page_texts = extract_page_texts(os.path.join(images_dir, pdf_files[0]))

                # Load image paths
                image_paths = {}
                for fname in sorted(os.listdir(images_dir)):
                    if fname.lower().endswith('.png'):
                        m = re.search(r'page_(\d+)\.png$', fname, re.IGNORECASE)
                        if m:
                            image_paths[int(m.group(1))] = os.path.join(images_dir, fname)
                print(f"[PAGE IMAGES] Loaded {len(image_paths)} page images, page_texts={'YES' if page_texts else 'NO'}")
            else:
                print("[PAGE IMAGES] No images folder provided — skipping image cues")
        except Exception as e:
            print(f"[PAGE IMAGES] WARNING: {e}")

        updated_pxmlcontent, processed_idxs_main, failed_inline_changes_main = process_with_llm(
            result_df=all_revisions_except_newline_insertions1, section_dict=section_dict,
            pxml_content=updated_pxmlcontent, client=client, failed_log_path=failed_log_path,
            match_threshold=0.5, model_name=model_name, chapter_number=chapter_number,
            page_texts=page_texts, image_paths=image_paths,
        )

        # Footnotes pass
        all_revisions_except_newline_insertions_footnotes = all_revisions_except_newline_insertions_new[
            (all_revisions_except_newline_insertions_new['text_before'] != '') &
            (all_revisions_except_newline_insertions_new['has_changes'] == True) &
            (all_revisions_except_newline_insertions_new['is_footnote_x'] == True)
        ]
        if 'parent_paragraph_idx_x' in all_revisions_except_newline_insertions_footnotes.columns:
            _parent_src = result_with_ba[['idx', 'type', 'text', 'text_before', 'text_after']].dropna(subset=['idx']).drop_duplicates(subset=['idx'], keep='first').copy()
            def _get_parent_text(row):
                rtype = str(row.get('type', '')).lower()
                if 'deletion' in rtype:
                    val = row.get('text_before', '') or ''
                    if str(val).strip():
                        return str(val)
                elif 'insertion' in rtype:
                    val = row.get('text_after', '') or ''
                    if str(val).strip():
                        return str(val)
                txt = row.get('text', '') or ''
                if str(txt).strip():
                    return str(txt)
                tb = row.get('text_before', '') or ''
                return str(tb) if str(tb).strip() else ''
            _parent_src['_parent_text'] = _parent_src.apply(_get_parent_text, axis=1)
            parent_para_text_map = _parent_src.set_index('idx')['_parent_text']
            all_revisions_except_newline_insertions_footnotes['parent_paragraph_text'] = all_revisions_except_newline_insertions_footnotes['parent_paragraph_idx_x'].map(parent_para_text_map).fillna('')
        else:
            all_revisions_except_newline_insertions_footnotes['parent_paragraph_text'] = ''

        all_revisions_except_newline_insertions_footnotes = fill_sections_backward_verbose(
            all_revisions_except_newline_insertions_footnotes, no_changes_df, inline_replacement_df, inline_insertion_df, chapter_number,
            verbose=True, client=client, model_name=model_name, designator_patterns=designator_patterns
        )

        updated_pxmlcontent, processed_idxs_footnotes, failed_inline_changes_footnotes = process_with_llm_footnotes(
            result_df=all_revisions_except_newline_insertions_footnotes, section_dict=section_dict,
            pxml_content=updated_pxmlcontent, client=client, failed_log_path=failed_log_path,
            match_threshold=0.5, model_name=model_name, chapter_number=chapter_number,
            page_texts=page_texts, image_paths=image_paths,
        )

        failed_inline_changes = (failed_inline_changes_main or []) + (failed_inline_changes_footnotes or [])

        # ═══ PIPELINE 6: Newline insertions ═══
        print("\n[PIPELINE 6] Newline paragraph insertions...")
        all_newline_insertions = result_with_ba_new[
            (result_with_ba_new['text_before'] == '') & (result_with_ba_new['has_changes'] == True) & (result_with_ba_new['is_footnote_x'] == False)
        ].drop_duplicates(subset=['idx'], keep='first')
        all_newline_insertions = all_newline_insertions[['idx', 'text', 'has_changes', 'type', 'paragraph_index']]
        all_newline_insertions = return_closest_previous_par(result_with_ba, all_newline_insertions)
        all_newline_insertions_sections = fill_sections_backward_verbose(
            all_newline_insertions, no_changes_df, inline_replacement_df, inline_insertion_df, chapter_number,
            verbose=True, client=client, model_name=model_name, designator_patterns=designator_patterns
        )

        section_dict, section_positions, pxml_content = parse_pxml_sections_from_str_dict_positions(updated_pxmlcontent)
        updated_pxmlcontent, failed_newline_insertions = process_all_newline_insertions_with_llm(
            pxml_content=pxml_content, section_dict=section_dict, section_positions=section_positions,
            all_newline_insertions=all_newline_insertions_sections, client=client, model_name=model_name,
            failed_log_path=os.path.join(output_dir, "failed_newline_insertions.txt"),
        )

        # ═══ PIPELINE 7: New footnote insertions ═══
        all_inserted_footnote_revisions = result_with_ba[
            (result_with_ba['is_footnote_x'] == True) & (result_with_ba['has_changes'] == True) & (result_with_ba['text_before'] == '')
        ].drop_duplicates(subset=['idx'], keep='first')
        if 'parent_paragraph_idx_x' in all_inserted_footnote_revisions.columns:
            _parent_src2 = result_with_ba[['idx', 'type', 'text', 'text_before', 'text_after']].dropna(subset=['idx']).drop_duplicates(subset=['idx'], keep='first').copy()
            def _get_parent_text2(row):
                rtype = str(row.get('type', '')).lower()
                if 'deletion' in rtype:
                    val = row.get('text_before', '') or ''
                    if str(val).strip():
                        return str(val)
                elif 'insertion' in rtype:
                    val = row.get('text_after', '') or ''
                    if str(val).strip():
                        return str(val)
                txt = row.get('text', '') or ''
                if str(txt).strip():
                    return str(txt)
                tb = row.get('text_before', '') or ''
                return str(tb) if str(tb).strip() else ''
            _parent_src2['_parent_text'] = _parent_src2.apply(_get_parent_text2, axis=1)
            parent_para_text_map = _parent_src2.set_index('idx')['_parent_text']
            all_inserted_footnote_revisions['full_para_where_footnote_is_present'] = all_inserted_footnote_revisions['parent_paragraph_idx_x'].map(parent_para_text_map).fillna('')
        else:
            all_inserted_footnote_revisions['full_para_where_footnote_is_present'] = ''

        section_dict, section_positions, pxml_content = parse_pxml_sections_from_str_dict_positions(updated_pxmlcontent)
        all_inserted_footnote_revisions_sections = fill_sections_backward_verbose(
            all_inserted_footnote_revisions, no_changes_df, inline_replacement_df, inline_insertion_df, chapter_number,
            verbose=True, client=client, model_name=model_name, designator_patterns=designator_patterns
        )
        updated_pxmlcontent, failed_footnote_insertions = process_all_footnote_insertions(
            all_inserted_footnote_revisions=all_inserted_footnote_revisions_sections,
            pxml_content=pxml_content, model_name=MODEL_OPUS,
            failed_log_path=os.path.join(output_dir, "failed_footnote_insertions.txt")
        )

        # Save before post-processing
        final_pxml_content = updated_pxmlcontent
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(final_pxml_content)

        # ═══ PIPELINE 8: Failed Revisions Report ═══
        print(f"\n{'='*80}\nWriting Consolidated Failed Revisions Excel...\n{'='*80}")
        consolidated_excel_path = os.path.join(output_dir, "Failed_revisions.xlsx")
        has_inline_failures = len(failed_inline_changes) > 0
        has_newline_failures = len(failed_newline_insertions) > 0
        has_footnote_failures = len(failed_footnote_insertions) > 0

        if has_inline_failures or has_newline_failures or has_footnote_failures:
            with pd.ExcelWriter(consolidated_excel_path, engine='openpyxl') as writer:
                from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

                header_font = Font(bold=True, color='FFFFFF', size=11)
                header_fill = PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid')
                header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                thin_border = Border(
                    left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin')
                )

                def apply_sheet_formatting(ws, df):
                    """Apply consistent formatting to a worksheet."""
                    for col_num in range(1, len(df.columns) + 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = header_alignment
                        cell.border = thin_border
                    for row_num in range(2, len(df) + 2):
                        for col_num in range(1, len(df.columns) + 1):
                            cell = ws.cell(row=row_num, column=col_num)
                            cell.alignment = cell_alignment
                            cell.border = thin_border
                    ws.row_dimensions[1].height = 25
                    for row_num in range(2, len(df) + 2):
                        ws.row_dimensions[row_num].height = 50
                    for col_num in range(1, len(df.columns) + 1):
                        col_letter = ws.cell(row=1, column=col_num).column_letter
                        max_len = max(len(str(ws.cell(row=1, column=col_num).value or '')),
                                      *[len(str(ws.cell(row=r, column=col_num).value or '')[:80]) for r in range(2, min(len(df)+2, 20))])
                        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)
                    ws.freeze_panes = 'A2'
                    ws.auto_filter.ref = ws.dimensions

                if has_inline_failures:
                    df_inline = pd.DataFrame(failed_inline_changes)
                    col_order = [c for c in ["Row Index", "Docx_Paragraph", "Section", "Designator", "Reason", "Change Type", "Inserted Text", "Deleted Text", "Previous Context", "Post Context"] if c in df_inline.columns]
                    df_inline[col_order].to_excel(writer, index=False, sheet_name='Inline Changes')
                    apply_sheet_formatting(writer.sheets['Inline Changes'], df_inline[col_order])
                if has_newline_failures:
                    df_nl = pd.DataFrame(failed_newline_insertions)
                    df_nl.to_excel(writer, index=False, sheet_name='Newline Insertions')
                    apply_sheet_formatting(writer.sheets['Newline Insertions'], df_nl)
                if has_footnote_failures:
                    df_fn = pd.DataFrame(failed_footnote_insertions)
                    df_fn.to_excel(writer, index=False, sheet_name='Footnote Insertions')
                    apply_sheet_formatting(writer.sheets['Footnote Insertions'], df_fn)

        # ═══ PIPELINE 9: Retry failed deletions ═══
        final_pxml_content, retry_success_count, retry_total, succeeded_indices, retry_rows_list = retry_failed_inline_deletions(
            failed_inline_changes=failed_inline_changes, pxml_content=final_pxml_content,
            client=client, model_name=model_name, chapter_number=chapter_number,
            page_texts=page_texts, image_paths=image_paths,
        )
        if retry_success_count > 0:
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write(final_pxml_content)

        # ═══ PIPELINE 10: Post Processing ═══
        print(f"\n{'='*80}\nPIPELINE-10: Post Processing (Final)\n{'='*80}")
        with open(os.path.join(output_dir, 'save_before_postprocessing.pxml'), 'w', encoding='utf-8') as f:
            f.write(final_pxml_content)
        final_pxml_content = postprocessing_fortags(final_pxml_content, product_type=PRODUCT_TYPE)
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(final_pxml_content)

        end_time = time.time()
        print(f"\n✅ Output saved to: {output_file}")
        print(f"Total runtime: {end_time - start_time:.2f} seconds")

        job.status = "completed"
        job.progress_pct = 100
        job.current_step = 10
        job.current_step_name = "Completed"
        job.completed_at = time.time()
        job.log = log_buffer.getvalue()[-50000:]
        job_store.update(job)

    except Exception as e:
        job.status = "failed"
        job.error = str(e)
        job.log = log_buffer.getvalue()[-50000:]
        job_store.update(job)
        print(f"\n❌ PIPELINE FAILED: {e}")

    finally:
        sys.stdout = sys.__stdout__
        sys.stderr = sys.__stderr__
        log_f.close()


# ─────────────────────────────────────────────
# API Endpoints
# ─────────────────────────────────────────────

@app.get("/api/health")
async def health_check():
    """Lightweight health endpoint for K8s probes."""
    return {"status": "healthy"}


@app.post("/api/process")
async def process_files(
    docx_file: UploadFile = File(...),
    pxml_file: UploadFile = File(...),
    images_zip: Optional[UploadFile] = File(None),
    product_type: str = Form("looseleaf"),
):
    """Upload DOCX + PXML + Images ZIP and start processing. Returns job_id for status polling."""
    import tempfile as _tempfile
    job_id = str(uuid.uuid4())[:8]
    
    # Use system temp folder for deployment compatibility (Kubernetes)
    base_temp = _tempfile.gettempdir()
    work_dir = os.path.join(base_temp, "pxml_pipeline_jobs", job_id)
    os.makedirs(work_dir, exist_ok=True)

    # Save uploaded files to temp
    docx_path = os.path.join(work_dir, docx_file.filename)
    pxml_path = os.path.join(work_dir, pxml_file.filename)
    with open(docx_path, "wb") as f:
        f.write(await docx_file.read())
    with open(pxml_path, "wb") as f:
        f.write(await pxml_file.read())

    # Extract images ZIP if provided
    images_dir = None
    if images_zip and images_zip.filename:
        images_dir = os.path.join(work_dir, "revision_images")
        os.makedirs(images_dir, exist_ok=True)
        zip_path = os.path.join(work_dir, images_zip.filename)
        with open(zip_path, "wb") as f:
            f.write(await images_zip.read())
        # Extract ZIP
        import zipfile as _zf
        with _zf.ZipFile(zip_path, 'r') as zf:
            zf.extractall(images_dir)
        os.remove(zip_path)
        # If ZIP contained a single subfolder, use that as images_dir
        extracted_items = os.listdir(images_dir)
        if len(extracted_items) == 1 and os.path.isdir(os.path.join(images_dir, extracted_items[0])):
            images_dir = os.path.join(images_dir, extracted_items[0])

    output_dir = os.path.join(work_dir, "output")
    os.makedirs(output_dir, exist_ok=True)

    job = JobStatus(job_id=job_id, output_dir=output_dir, file_name=docx_file.filename)
    job_store.create(job)

    # Dispatch to Celery worker
    run_pipeline_task.delay(job_id, docx_path, pxml_path, output_dir, product_type, images_dir)

    return {"job_id": job_id, "status": "queued"}


@app.post("/api/process-batch")
async def process_batch(
    files: list[UploadFile] = File(...),
    product_type: str = Form("looseleaf"),
):
    """Upload a zip of folders (each with DOCX + PXML). Returns batch job IDs."""
    batch_id = str(uuid.uuid4())[:8]
    base_dir = os.path.dirname(os.path.abspath(__file__))
    batch_dir = os.path.join(base_dir, "pipeline_jobs", f"batch_{batch_id}")
    os.makedirs(batch_dir, exist_ok=True)

    job_ids = []

    for uploaded_file in files:
        if uploaded_file.filename.endswith(".zip"):
            # Extract zip
            zip_path = os.path.join(batch_dir, uploaded_file.filename)
            with open(zip_path, "wb") as f:
                f.write(await uploaded_file.read())
            import zipfile
            with zipfile.ZipFile(zip_path, 'r') as zf:
                zf.extractall(batch_dir)
            os.remove(zip_path)

    # Find all subfolders with DOCX + PXML pairs
    for folder in Path(batch_dir).iterdir():
        if not folder.is_dir():
            continue
        docx_files = list(folder.glob("*.docx")) + list(folder.glob("*.doc"))
        pxml_files = list(folder.glob("*.pxml"))
        if docx_files and pxml_files:
            job_id = str(uuid.uuid4())[:8]
            output_dir = str(folder / "output")
            os.makedirs(output_dir, exist_ok=True)
            job = JobStatus(job_id=job_id, output_dir=output_dir, file_name=docx_files[0].name)
            job_store.create(job)
            run_pipeline_task.delay(job_id, str(docx_files[0]), str(pxml_files[0]), output_dir, product_type)
            job_ids.append(job_id)

    return {"batch_id": batch_id, "job_ids": job_ids}


@app.get("/api/status/{job_id}")
async def get_status(job_id: str):
    """Poll job status."""
    job = job_store.get(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    return job.model_dump()


@app.get("/api/download/{job_id}")
async def download_output(job_id: str):
    """Download output folder as zip."""
    job = job_store.get(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    if job.status != "completed":
        raise HTTPException(400, f"Job not complete. Status: {job.status}")

    base_dir = os.path.dirname(os.path.abspath(__file__))
    zip_path = os.path.join(base_dir, "pipeline_jobs", f"{job_id}_output.zip")
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(job.output_dir):
            for file in files:
                filepath = os.path.join(root, file)
                arcname = os.path.relpath(filepath, job.output_dir)
                zf.write(filepath, arcname)
    return FileResponse(zip_path, filename=f"{job_id}_output.zip", media_type="application/zip")


@app.post("/api/validate/{job_id}")
async def run_validation(job_id: str):
    """Run validation on a completed job."""
    job = job_store.get(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    if job.status != "completed":
        raise HTTPException(400, "Pipeline must complete before validation")

    try:
        from exe_track_changes_1 import run_validation_for_gui
        script_dir = os.path.dirname(os.path.abspath(__file__))
        prompts_dir = os.path.join(script_dir, "validation_prompts")

        # Ensure poppler is findable (add to PATH if exists in project root or local installs)
        poppler_candidates = [
            os.path.join(script_dir, "poppler", "Library", "bin"),
            os.path.join(script_dir, "poppler", "bin"),
            os.path.join(script_dir, "poppler"),
            r"C:\Program Files\poppler-25.07.0\Library\bin",
            r"C:\Program Files\poppler-25.07.0\bin",
            r"C:\Program Files\poppler-25.07.0",
            r"C:\poppler\Library\bin",
            r"C:\poppler\bin",
        ]
        for poppler_bin in poppler_candidates:
            if os.path.isdir(poppler_bin) and poppler_bin not in os.environ.get("PATH", ""):
                os.environ["PATH"] = poppler_bin + os.pathsep + os.environ.get("PATH", "")
                break

        # Find the DOCX in the job folder
        job_dir = os.path.dirname(job.output_dir)
        docx_files = list(Path(job_dir).glob("*.docx")) + list(Path(job_dir).glob("*.doc"))
        if not docx_files:
            raise HTTPException(400, "DOCX not found in job folder")

        # Locate the uploaded images folder (extracted from ZIP during upload)
        images_folder = None
        _revision_images_dir = os.path.join(job_dir, "revision_images")
        if os.path.isdir(_revision_images_dir):
            # If ZIP contained a single subfolder, use that
            _items = os.listdir(_revision_images_dir)
            if len(_items) == 1 and os.path.isdir(os.path.join(_revision_images_dir, _items[0])):
                images_folder = os.path.join(_revision_images_dir, _items[0])
            else:
                images_folder = _revision_images_dir

        excel_path = run_validation_for_gui(
            docx_path=str(docx_files[0]),
            final_xml_path=job.output_file,
            output_dir=job.output_dir,
            prompt_path=os.path.join(prompts_dir, "prompt.txt"),
            validation_prompt_path=os.path.join(prompts_dir, "prompt_validation.txt"),
            images_folder=images_folder,
        )
        return {"status": "completed", "excel_path": excel_path}
    except Exception as e:
        raise HTTPException(500, f"Validation failed: {e}")


@app.get("/api/download-validation/{job_id}")
async def download_validation(job_id: str):
    """Download validation Excel file."""
    job = job_store.get(job_id)
    if not job:
        raise HTTPException(404, "Job not found")

    # Find validation Excel in output dir
    excel_files = list(Path(job.output_dir).glob("final_validation_report_*.xlsx"))
    if not excel_files:
        raise HTTPException(404, "Validation report not found. Run validation first.")
    return FileResponse(str(excel_files[0]), filename=excel_files[0].name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/log/{job_id}")
async def get_log(job_id: str):
    """Get live log output."""
    job = job_store.get(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    return {"log": job.log}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
